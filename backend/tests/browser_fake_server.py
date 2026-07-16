"""Start an isolated backend with a local-only PPT generator for browser checks."""

from __future__ import annotations

import os
import time
import uuid
from pathlib import Path


os.environ["EDUAGENT_SKIP_ENV_FILE"] = "1"
os.environ.setdefault("LLM_PROVIDER", "mock")
os.environ.setdefault("RAG_ENABLED", "false")
if not os.environ.get("DATABASE_URL"):
    raise SystemExit("DATABASE_URL must point to a temporary SQLite file")

from pptx import Presentation  # noqa: E402

from app.config import settings  # noqa: E402
from app.services import ppt_generator  # noqa: E402
from app.routers import chat_router  # noqa: E402
from app.routers import product  # noqa: E402
from app.services.conversation_state import conversation_store  # noqa: E402
from app.services import section_resource_recommendations  # noqa: E402
from app.services.search_client import SearchError, SearchResponse, SearchResultItem  # noqa: E402


def _fake_generate_pptx(topic: str, difficulty: str = "medium", session_id: str = "") -> str:
    output_dir = Path(settings.project_root) / "outputs" / "ppt"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"browser-{uuid.uuid4().hex}.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = topic
    slide.placeholders[1].text = f"Local browser fixture ({difficulty})"
    presentation.save(path)
    return str(path)


ppt_generator.generate_pptx = _fake_generate_pptx


async def _fake_chat_pipeline(**state: object) -> dict[str, object]:
    """Deterministic local-only replies for browser profile-boundary checks."""
    message = str(state.get("user_message") or "")
    profile = state.get("profile_v2") if isinstance(state.get("profile_v2"), dict) else {}
    context = profile.get("subject_context") if isinstance(profile.get("subject_context"), dict) else {}
    if "年级" in message:
        reply = f"你目前是{context.get('background') or '未记录'}。"
    elif "偏好" in message:
        preferences = "、".join(context.get("resource_preferences") or [])
        reply = f"你偏好通过{preferences or '未记录'}学习。"
    else:
        reply = "已记录当前对话中的明确学习信息。"
    return {"final_reply": reply, "agents_run": ["fake_conversation"], "pipeline_executed": True}


chat_router.run_pipeline = _fake_chat_pipeline
conversation_store.extract_facts_with_llm = lambda state, message: conversation_store.extract_facts(state, message)

_real_general_resource_generation = product._generate_general_resource


def _delayed_general_resource_generation(*args: object, **kwargs: object) -> dict[str, object]:
    request = args[0] if args else kwargs.get("request")
    if isinstance(request, dict) and request.get("resourceType") == "manim":
        raise RuntimeError("provider_not_configured")
    delay_ms = int(os.getenv("EDUAGENT_GENERAL_RESOURCE_DELAY_MS", "0") or 0)
    if delay_ms:
        time.sleep(delay_ms / 1000)
    return _real_general_resource_generation(*args, **kwargs)


product._generate_general_resource = _delayed_general_resource_generation


class _LocalSearchClient:
    """Deterministic public-looking URLs for isolated browser search checks."""

    def __init__(self) -> None:
        self._unavailable = False

    def search(self, query: str, max_results: int = 5) -> SearchResponse:
        if "联网故障" in query:
            self._unavailable = True
        if self._unavailable:
            raise SearchError("fixture search unavailable")
        lowered = query.lower()
        if "视频" in query or "video" in lowered:
            item = SearchResultItem("递归调用栈教学视频", "https://www.bilibili.com/video/BV1fixture", "递归调用栈、栈帧与返回地址讲解", "fixture")
        elif "课程" in query or "mooc" in lowered:
            item = SearchResultItem("递归调用栈公开课程", "https://www.icourse163.org/learn/DS-100", "数据结构课程中的递归调用栈", "fixture")
        elif "课件" in query or "讲义" in query or "lecture notes" in lowered:
            item = SearchResultItem("递归调用栈讲义", "https://cs.fixture.edu/notes/recursion-call-stack.pdf", "递归调用栈和栈帧讲义", "fixture")
        elif "paper" in lowered or "research" in lowered:
            item = SearchResultItem("Recursion Call Stack Study", "https://arxiv.org/abs/2401.12345", "recursion call stack stack frame analysis", "fixture")
        else:
            item = SearchResultItem("递归调用栈学习文章", "https://learn.fixture.edu/recursion-call-stack", "递归调用栈、栈帧与返回地址", "fixture")
        return SearchResponse(query=query, results=[item], total_estimated=1, source="fixture")


section_resource_recommendations.get_search_client = lambda _name: _LocalSearchClient()
section_resource_recommendations.SectionResourceRecommendationService._auto_ingest = lambda self, _resources: None
def _unavailable_paper_search(*_args: object, **_kwargs: object) -> SearchResponse:
    raise SearchError("fixture paper search unavailable")


section_resource_recommendations.search_crossref = _unavailable_paper_search
section_resource_recommendations.search_arxiv = _unavailable_paper_search
_real_recommend = section_resource_recommendations.SectionResourceRecommendationService.recommend


def _delayed_recommend(*args: object, **kwargs: object) -> dict[str, object]:
    delay_ms = int(os.getenv("EDUAGENT_RESOURCE_SEARCH_DELAY_MS", "0") or 0)
    if delay_ms:
        time.sleep(delay_ms / 1000)
    return _real_recommend(*args, **kwargs)


section_resource_recommendations.SectionResourceRecommendationService.recommend = _delayed_recommend

from app.main import app  # noqa: E402


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=int(os.getenv("PORT", "8010")))
