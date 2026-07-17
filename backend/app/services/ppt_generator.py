"""Generate .pptx files with rich layouts via DeepSeek + python-pptx.

Architecture:
  generate_pptx(topic, difficulty, session_id)
    ├─ generate_slide_outline()  →  structured JSON outline via DeepSeek
    └─ build_pptx()              →  physical .pptx via python-pptx
       return (filepath, outline)

Each slide has a `style` field chosen by the LLM:
  "cover", "concept", "definition", "content", "comparison",
  "process", "example", "summary", "qa"
Each style uses a different visual layout.
"""

from __future__ import annotations

import json
import logging
import os
import uuid
from pathlib import Path
from typing import Any

import requests

logger = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).resolve().parents[3] / "outputs" / "ppt"

try:
    from pptx import Presentation as _PptxPresentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    _PPTX_OK = True
except ImportError:
    _PPTX_OK = False

# ── Color palette ────────────────────────────────────────────────────────

C = {
    "bg_dark":    (0x1A, 0x1A, 0x2E),
    "bg_cover":   (0x0F, 0x0F, 0x1E),
    "bg_even":    (0x1E, 0x1E, 0x34),
    "bg_odd":     (0x16, 0x16, 0x2A),
    "bg_sum":     (0x0A, 0x0A, 0x18),
    "accent":     (0x3B, 0x82, 0xF6),   # blue
    "accent2":    (0x8B, 0x5C, 0xF6),   # purple
    "accent3":    (0x06, 0xD6, 0xA0),   # green
    "gold":       (0xF5, 0x9E, 0x0B),
    "white":      (0xF1, 0xF5, 0xF9),
    "gray":       (0xA0, 0xA8, 0xB8),
    "gray_dark":  (0x6B, 0x72, 0x80),
    "box_bg":     (0x25, 0x25, 0x3D),   # slightly lighter for callout boxes
    "highlight":  (0xE8, 0x6C, 0x00),   # orange highlight
}


def _rgb(name):
    return RGBColor(*C[name])


def _tb(slide, left, top, w, h):
    return slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h)).text_frame


def _p(tf, text, size=18, color="white", bold=False, align=None, space_after=6, italic=False):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = _rgb(color)
    p.font.bold = bold
    p.font.italic = italic
    if align:
        p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    return p


def _box(slide, left, top, w, h, color="box_bg", radius=True):
    """Add a rounded rectangle callout box."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    # Adjust corner radius
    if radius:
        shape.adjustments[0] = 0.1
    return shape.text_frame


# ── Image generation (Seedream) ──────────────────────────────────────────

_IMAGE_DIR = OUTPUT_DIR / "images"
_MAX_IMAGE_SLIDES = 2


def _gen_images(outline):
    slides = [(i, s["image_prompt"]) for i, s in enumerate(outline) if s.get("image_prompt")][:_MAX_IMAGE_SLIDES]
    if not slides:
        return {}
    _IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    try:
        from app.services.multimodal_provider import SeedreamImageProvider
        provider = SeedreamImageProvider()
    except Exception:
        return {}

    results = {}
    for idx, prompt in slides:
        try:
            resp = provider.run({"prompt": prompt})
            if resp.get("status") != "success":
                continue
            urls = resp.get("result", {}).get("image_urls", [])
            if not urls:
                continue
            clean = urls[0].split("?", 1)[0]
            ext = clean.rsplit(".", 1)[-1] if "." in clean else "png"
            path = str(_IMAGE_DIR / f"s{idx}_{uuid.uuid4().hex[:8]}.{ext}")
            r = requests.get(urls[0], timeout=60)
            if r.status_code == 200:
                with open(path, "wb") as f:
                    f.write(r.content)
                results[idx] = path
                logger.info("PPT image %d saved: %s", idx, path)
        except Exception as e:
            logger.warning("PPT image %d failed: %s", idx, e)
    return results


# ── Step 1: Outline ──────────────────────────────────────────────────────


def _llm(messages, temp=0.4, max_tk=8000):
    try:
        from app.services.llm_client import get_llm_client
        llm = get_llm_client()
        if llm and llm.is_available():
            return llm.chat(messages=messages, temperature=temp, max_tokens=max_tk)
    except Exception as e:
        logger.warning("PPT LLM: %s", e)
    return None


def generate_slide_outline(topic: str, difficulty: str) -> list[dict[str, Any]] | None:
    """DeepSeek generates a rich slide deck outline.

    Returns a list of slides. Each slide has:
      title, bullets (list), content (detail paragraph),
      style (see below), notes, image_prompt (optional).

    Styles: cover | concept | definition | content | comparison
            | process | example | summary | qa
    """
    prompt = f"""你是一位资深课程设计师，为「{topic}」设计 PPT 演示文稿（难度：{difficulty}）。

输出 JSON 数组，每页格式：
{{
  "title": "标题",
  "bullets": ["分点1", "分点2", ...],     ← 自由发挥，不限数量
  "content": "详细讲解段落，用通俗语言+例子讲清楚",  ← 150-300字
  "style": "版式类型",
  "notes": "讲师备注（简短）",
  "image_prompt": "配图描述（仅封面和关键概念页）"
}}

版式类型（请根据内容选择最合适的）：
  cover      封面页（必须有 image_prompt）
  concept    概念引入——左侧配图/右侧关键点
  definition 定义讲解——突出核心定义+解释
  content    标准内容——标题+要点+详细段落
  comparison 对比——两栏对比/相似vs不同
  process    流程步骤——分步骤逐步讲解
  example    实例应用——案例+分析
  summary    总结回顾——结构化梳理
  qa         思考题——开放问题

页数：由内容深度决定，不限页数，讲清楚为止（建议至少10页）。
每页 bullets 不限条数。content 一定要有实质内容，用例子/类比讲明白。
封面页必须提供 image_prompt。

只输出 JSON，不要其他文字。"""

    raw = _llm([
        {"role": "system", "content": "你是一个严格的 JSON 输出器。输出合法的 JSON 数组，不要 markdown 包裹。"},
        {"role": "user", "content": prompt},
    ], max_tk=8000)
    if not raw:
        return None

    text = raw.strip()
    if "```" in text:
        parts = text.split("```")
        for i, p in enumerate(parts):
            if p.strip().startswith("json"):
                text = p.strip()[4:].strip()
                break
            if i % 2 == 1:
                text = p.strip()
                break

    try:
        outline = json.loads(text)
    except json.JSONDecodeError:
        logger.warning("PPT outline: JSON parse failed")
        return None

    if not isinstance(outline, list) or len(outline) < 3:
        return None

    VALID_STYLES = {"cover", "concept", "definition", "content", "comparison", "process", "example", "summary", "qa"}
    validated = []
    for s in outline:
        if not isinstance(s, dict):
            continue
        title = str(s.get("title", "")).strip()
        if not title:
            continue
        style = str(s.get("style", "content")).strip()
        if style not in VALID_STYLES:
            style = "content"
        entry = {"title": title, "style": style}
        bullets = s.get("bullets", [])
        entry["bullets"] = [str(b).strip() for b in bullets if str(b).strip()] if isinstance(bullets, list) else [str(bullets)]
        content = str(s.get("content", "")).strip()
        if content:
            entry["content"] = content
        notes = str(s.get("notes", "")).strip()
        if notes:
            entry["notes"] = notes
        ip = str(s.get("image_prompt", "")).strip()
        if ip:
            entry["image_prompt"] = ip
        validated.append(entry)

    if len(validated) < 3:
        return None
    logger.info("PPT outline: %d slides for %s", len(validated), topic)
    return validated


# ── Step 2: Build layouts ────────────────────────────────────────────────


def _set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _bar(slide, l, t, w, h, color, shape_type=MSO_SHAPE.RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    return shape


def _img(slide, path, l, t, w, h):
    if os.path.exists(path):
        slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))
        return True
    return False


def _render_cover(slide, data, topic, img_path):
    """Cover with title + optional image."""
    _set_bg(slide, "bg_cover")
    _bar(slide, 0, 0, 13.33, 3.2, "accent3")
    _bar(slide, 0, 0, 13.33, 0.06, "accent")

    if img_path:
        _img(slide, img_path, 7.5, 0.8, 5.3, 6.0)

    text_w = 6.8 if img_path else 11.5
    tf = _tb(slide, 0.8 if img_path else 1.2, 1.0, text_w, 1.5)
    _p(tf, data["title"], size=48, color="gold", bold=True, space_after=12)
    _p(tf, topic, size=24, color="gray", space_after=0)

    bulk = data.get("content", "")
    if bulk:
        tf2 = _tb(slide, 0.8, 3.5, text_w, 2.5)
        _p(tf2, bulk, size=15, color="gray_dark", space_after=0)

    note = data.get("notes", "")
    if note:
        tf3 = _tb(slide, 0.8, 6.3, text_w, 0.6)
        _p(tf3, f"💡 {note}", size=13, color="accent", space_after=0)


def _render_concept(slide, data, idx, total, img_path):
    """Concept intro: left image OR left text + right bullets."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "accent2")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    if img_path:
        _img(slide, img_path, 0.5, 1.8, 5.5, 5.0)

    text_left = 6.5 if img_path else 0.8
    tf = _tb(slide, text_left, 0.5, 6.0, 0.7)
    _p(tf, data["title"], size=36, color="white", bold=True, space_after=0)

    for bi, b in enumerate(data.get("bullets", [])):
        c = ["accent", "accent2", "accent3", "gold", "accent", "accent2"][bi % 6]
        tf2 = _tb(slide, text_left + 0.3, 1.6 + bi * 0.65, 5.5, 0.6)
        _p(tf2, f"▸ {b}", size=18, color=c, space_after=0)

    bulk = data.get("content", "")
    if bulk:
        y = 1.6 + len(data.get("bullets", [])) * 0.65 + 0.2
        _bar(slide, text_left + 0.3, y - 0.1, 5.5, 0.015, "gray_dark")
        tf3 = _tb(slide, text_left + 0.3, y + 0.1, 5.5, 1.5)
        _p(tf3, bulk, size=14, color="gray", space_after=0)

    note = data.get("notes", "")
    if note:
        tf4 = _tb(slide, text_left + 0.3, 6.5, 5.5, 0.5)
        _p(tf4, f"💡 {note}", size=12, color="accent2", italic=True, space_after=0)


def _render_definition(slide, data, idx, total):
    """Definition: callout box + bullets + detail."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "gold")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf = _tb(slide, 0.8, 0.4, 11.0, 0.7)
    _p(tf, data["title"], size=36, color="gold", bold=True, space_after=0, align=PP_ALIGN.LEFT)

    # Definition callout box
    bulk = data.get("content", "")
    if bulk:
        box_tf = _box(slide, 0.8, 1.4, 11.5, 1.8, "box_bg")
        _p(box_tf, "📖 核心定义", size=13, color="gold", bold=True, space_after=6)
        _p(box_tf, bulk, size=16, color="white", space_after=0)

    # Bullets below
    bullets = data.get("bullets", [])
    for bi, b in enumerate(bullets[:4]):
        c = ["accent", "accent3", "accent2", "gold"][bi % 4]
        tf2 = _tb(slide, 1.0, 3.6 + bi * 0.65, 11.0, 0.6)
        _p(tf2, f"▪ {b}", size=18, color=c, space_after=0)

    note = data.get("notes", "")
    if note:
        tf4 = _tb(slide, 0.8, 6.5, 11.0, 0.5)
        _p(tf4, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


def _render_content(slide, data, idx, total, img_path):
    """Standard content: title + bullets + detail paragraph."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "accent")
    _bar(slide, 12.6, 0.4, 0.12, 0.12, ["accent", "accent2", "accent3"][idx % 3])

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    text_w = 7.5 if img_path else 11.5
    if img_path:
        _img(slide, img_path, 8.0, 1.4, 4.8, 5.0)

    tf = _tb(slide, 0.8, 0.4, text_w, 0.7)
    _p(tf, data["title"], size=34, color="white", bold=True, space_after=0)

    bullets = data.get("bullets", [])
    for bi, b in enumerate(bullets):
        c = ["accent", "accent2", "accent3", "gold", "accent", "accent2"][bi % 6]
        tf2 = _tb(slide, 1.0, 1.5 + bi * 0.55, text_w - 0.3, 0.5)
        _p(tf2, f"{['▸', '▹', '◆', '▪', '•'][bi % 5]}  {b}", size=18, color=c, space_after=0)

    bulk = data.get("content", "")
    if bulk:
        y = 1.5 + len(bullets) * 0.55 + 0.2 if bullets else 2.0
        _bar(slide, 0.8, y - 0.08, text_w, 0.015, "gray_dark")
        tf3 = _tb(slide, 0.8, y + 0.1, text_w, 1.8)
        _p(tf3, bulk, size=14, color="gray", space_after=0)

    note = data.get("notes", "")
    if note:
        tf4 = _tb(slide, 0.8, 6.5, text_w, 0.5)
        _p(tf4, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


def _render_comparison(slide, data, idx, total):
    """Two-column comparison with divider."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "accent3")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf = _tb(slide, 0.8, 0.4, 11.0, 0.7)
    _p(tf, data["title"], size=34, color="accent3", bold=True, space_after=0)

    # Vertical divider
    _bar(slide, 6.5, 1.6, 0.03, 5.0, "gray_dark")

    bullets = data.get("bullets", [])
    mid = len(bullets) // 2
    for bi, b in enumerate(bullets[:mid]):
        side = "left"
        tf2 = _tb(slide, 0.8, 1.8 + bi * 0.75, 5.3, 0.7)
        _p(tf2, f"▸ {b}", size=17, color=["accent", "accent2"][bi % 2], space_after=0)
    for bi, b in enumerate(bullets[mid:]):
        tf2 = _tb(slide, 6.8, 1.8 + bi * 0.75, 5.3, 0.7)
        _p(tf2, f"▸ {b}", size=17, color=["accent3", "gold"][bi % 2], space_after=0)

    bulk = data.get("content", "")
    if bulk:
        tf3 = _tb(slide, 0.8, 5.5, 11.5, 1.0)
        _p(tf3, bulk, size=13, color="gray", space_after=0)

    note = data.get("notes", "")
    if note:
        tf4 = _tb(slide, 0.8, 6.5, 11.0, 0.5)
        _p(tf4, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


def _render_process(slide, data, idx, total):
    """Step-by-step with numbered boxes."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "accent2")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf = _tb(slide, 0.8, 0.4, 11.0, 0.7)
    _p(tf, data["title"], size=34, color="accent2", bold=True, space_after=0)

    bullets = data.get("bullets", [])
    cols = min(3, max(1, len(bullets) // 2))
    box_w = (11.5 - (cols - 1) * 0.3) / cols

    for bi, b in enumerate(bullets):
        row = bi // cols
        col = bi % cols
        l = 0.8 + col * (box_w + 0.3)
        t = 1.8 + row * 1.8
        box_tf = _box(slide, l, t, box_w, 1.5, "box_bg")
        _p(box_tf, f"Step {bi + 1}", size=11, color="gold", bold=True, space_after=4)
        _p(box_tf, b, size=14, color="white", space_after=0)

        # Arrow between columns
        if col < cols - 1 and len(bullets) > bi + 1:
            _bar(slide, l + box_w + 0.05, t + 0.6, 0.2, 0.04, "accent3")

    bulk = data.get("content", "")
    if bulk:
        y = 1.8 + ((len(bullets) - 1) // cols + 1) * 1.8 + 0.1
        tf2 = _tb(slide, 0.8, y, 11.5, 1.0)
        _p(tf2, bulk, size=14, color="gray", space_after=0)

    note = data.get("notes", "")
    if note:
        tf3 = _tb(slide, 0.8, 6.5, 11.0, 0.5)
        _p(tf3, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


def _render_example(slide, data, idx, total):
    """Example: highlight box + analysis."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "accent3")
    _bar(slide, 12.6, 0.4, 0.12, 0.12, "accent3")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf = _tb(slide, 0.8, 0.4, 11.0, 0.7)
    _p(tf, data["title"], size=34, color="accent3", bold=True, space_after=0)

    # Example callout box (top highlight)
    bulk = data.get("content", "")
    if bulk:
        box_tf = _box(slide, 0.8, 1.4, 11.5, 2.5, "box_bg")
        _p(box_tf, "💡 实例分析", size=13, color="gold", bold=True, space_after=6)
        _p(box_tf, bulk, size=16, color="white", space_after=0)

    bullets = data.get("bullets", [])
    for bi, b in enumerate(bullets):
        c = ["accent", "accent2", "accent3", "gold"][bi % 4]
        tf2 = _tb(slide, 1.0, 4.2 + bi * 0.55, 11.0, 0.5)
        _p(tf2, f"✓ {b}", size=17, color=c, space_after=0)

    note = data.get("notes", "")
    if note:
        tf3 = _tb(slide, 0.8, 6.5, 11.0, 0.5)
        _p(tf3, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


def _render_summary(slide, data, idx, total):
    """Summary: structured recap."""
    _set_bg(slide, "bg_sum")
    _bar(slide, 0, 0, 13.33, 0.08, "gold")
    _bar(slide, 0, 7.4, 13.33, 0.1, "accent3")
    _bar(slide, 0, 0, 0.4, 7.5, "accent3")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf_label = _tb(slide, 1.2, 0.5, 4.0, 0.5)
    _p(tf_label, "📌 课程总结", size=16, color="gold", bold=True, space_after=0)
    tf = _tb(slide, 1.2, 1.2, 10.5, 0.8)
    _p(tf, data["title"], size=40, color="gold", bold=True, space_after=12)

    for bi, b in enumerate(data.get("bullets", [])):
        c = ["accent", "accent2", "accent3", "gold"][bi % 4]
        tf2 = _tb(slide, 1.2, 2.5 + bi * 0.65, 10.5, 0.6)
        _p(tf2, f"✦  {b}", size=19, color=c, space_after=0)

    bulk = data.get("content", "")
    if bulk:
        tf3 = _tb(slide, 1.2, 5.8, 10.5, 1.0)
        _p(tf3, bulk, size=14, color="gray", space_after=0)

    note = data.get("notes", "")
    if note:
        tf4 = _tb(slide, 1.2, 6.5, 10.5, 0.5)
        _p(tf4, f"💡 {note}", size=12, color="accent3", italic=True, space_after=0)


def _render_qa(slide, data, idx, total):
    """Q&A: questions in callout boxes."""
    _set_bg(slide, "bg_even" if idx % 2 == 0 else "bg_odd")
    _bar(slide, 0.8, 1.3, 3.0, 0.04, "gold")

    tf_page = _tb(slide, 12.0, 0.3, 1.0, 0.5)
    _p(tf_page, f"{idx + 1}/{total}", size=11, color="gray_dark", align=PP_ALIGN.RIGHT, space_after=0)

    tf = _tb(slide, 0.8, 0.4, 11.0, 0.7)
    _p(tf, data["title"], size=34, color="gold", bold=True, space_after=0)

    bullets = data.get("bullets", [])
    for bi, b in enumerate(bullets):
        y = 1.5 + bi * 1.3
        box_tf = _box(slide, 0.8, y, 11.5, 1.0, "box_bg")
        _p(box_tf, f"❓ 思考 {bi + 1}", size=12, color="gold", bold=True, space_after=4)
        _p(box_tf, b, size=16, color="white", space_after=0)

    note = data.get("notes", "")
    if note:
        tf2 = _tb(slide, 0.8, 6.5, 11.0, 0.5)
        _p(tf2, f"💡 {note}", size=12, color="accent", italic=True, space_after=0)


_RENDERERS = {
    "cover":       _render_cover,
    "concept":     _render_concept,
    "definition": _render_definition,
    "content":    _render_content,
    "comparison": _render_comparison,
    "process":    _render_process,
    "example":    _render_example,
    "summary":    _render_summary,
    "qa":         _render_qa,
}


def build_pptx(outline, topic, images=None):
    if not _PPTX_OK:
        logger.warning("python-pptx not installed")
        return None

    try:
        prs = _PptxPresentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        total = len(outline)

        for idx, data in enumerate(outline):
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
            style = data.get("style", "content")
            img_path = images.get(idx) if images else None

            renderer = _RENDERERS.get(style, _render_content)
            try:
                if style == "cover":
                    renderer(slide, data, topic, img_path)
                elif style == "concept":
                    _render_concept(slide, data, idx, total, img_path)
                elif style == "definition":
                    _render_definition(slide, data, idx, total)
                elif style == "comparison":
                    _render_comparison(slide, data, idx, total)
                elif style == "process":
                    _render_process(slide, data, idx, total)
                elif style == "example":
                    _render_example(slide, data, idx, total)
                elif style == "summary":
                    _render_summary(slide, data, idx, total)
                elif style == "qa":
                    _render_qa(slide, data, idx, total)
                else:
                    _render_content(slide, data, idx, total, img_path)
            except Exception as e:
                logger.warning("PPT slide %d render error (%s): %s", idx, style, e)
                _render_content(slide, data, idx, total, img_path)

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        path = str(OUTPUT_DIR / f"{uuid.uuid4().hex}.pptx")
        prs.save(path)
        logger.info("PPTX: %s (%d slides)", path, total)
        return path
    except Exception as e:
        logger.warning("PPT build failed: %s", e)
        return None


# ── Entry point ──────────────────────────────────────────────────────────


def generate_pptx(topic: str, difficulty: str = "medium", session_id: str = "") -> tuple[str | None, list[dict[str, Any]] | None]:
    logger.info("PPT gen: topic=%s difficulty=%s", topic, difficulty)

    outline = generate_slide_outline(topic, difficulty)
    if not outline:
        logger.warning("PPT gen aborted: outline failed")
        return None, None

    images = _gen_images(outline)
    if images:
        logger.info("PPT images: %d slides", len(images))

    path = build_pptx(outline, topic, images)
    if not path:
        return None, None

    logger.info("PPT complete: %s", path)
    return path, outline
